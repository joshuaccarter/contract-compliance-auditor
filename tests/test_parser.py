"""Unit tests for the contract document parser and structural clause chunker."""

from pathlib import Path

import pytest
from docx import Document

from src.parsing.document_parser import ContractClause, DocumentParser, ParsedContract


def test_parser_markdown_normalization():
    """Verifies parsing, section extraction, and clause chunking on markdown contracts."""
    contract_file = Path("data/raw_contracts/generic_tech_vendor_contract.md")
    assert contract_file.exists(), "Test contract fixture missing"

    parser = DocumentParser()
    result = parser.parse(contract_file)

    assert isinstance(result, ParsedContract)
    assert result.file_name == "generic_tech_vendor_contract.md"
    assert result.file_type == ".md"
    assert result.estimated_tokens > 100
    assert len(result.detected_sections) >= 4
    assert len(result.clauses) >= 10

    # Verify key sections were detected by heading heuristics
    joined_sections = " ".join(result.detected_sections).lower()
    assert "fees" in joined_sections or "invoicing" in joined_sections
    assert "term" in joined_sections or "renewals" in joined_sections


def test_parser_clause_chunking_and_cross_references():
    """Verifies that numbered sub-clauses, headings, and cross-references are extracted accurately."""
    contract_file = Path("data/raw_contracts/generic_tech_vendor_contract.md")
    parser = DocumentParser()
    parsed = parser.parse(contract_file)

    clauses_by_id = {c.clause_id: c for c in parsed.clauses}

    assert "PREAMBLE" in clauses_by_id
    assert "Section 3.2" in clauses_by_id
    assert "Section 5.1" in clauses_by_id
    assert "Section 8.1" in clauses_by_id
    assert "Section 11.2" in clauses_by_id
    assert "Section 14.1" in clauses_by_id

    # Check parent section heading association and clause title
    sec_3_2 = clauses_by_id["Section 3.2"]
    assert isinstance(sec_3_2, ContractClause)
    assert "Section 3" in sec_3_2.section_heading
    assert sec_3_2.clause_title == "Invoicing Schedule"
    assert "thirty (30) calendar days" in sec_3_2.text

    # Check cross-reference detection in Section 11.2 ("UNDER SECTION 9")
    sec_11_2 = clauses_by_id["Section 11.2"]
    assert "9" in sec_11_2.referenced_sections


def test_parser_oversized_clause_splitting():
    """Verifies that long clauses exceeding max_clause_chars are split with sentence overlap."""
    long_clause = (
        "## Section 4: Data Security\n"
        "4.1 Security Controls. "
        + " ".join(
            f"Sentence number {i} establishes mandatory encryption and audit logging controls."
            for i in range(1, 25)
        )
    )
    parser = DocumentParser(max_clause_chars=300, overlap_chars=80)
    parsed = parser.parse_text(long_clause)

    part_ids = [c.clause_id for c in parsed.clauses]
    assert len(parsed.clauses) >= 4
    assert part_ids[0] == "Section 4.1 (Part 1)"
    assert part_ids[1] == "Section 4.1 (Part 2)"
    # Verify sentence overlap is preserved across consecutive parts
    assert parsed.clauses[0].text.split(".")[-2].strip() in parsed.clauses[1].text


def test_parser_docx_reading_order_preservation(tmp_path):
    """Verifies that inline tables in DOCX files remain in reading order between paragraphs."""
    docx_path = tmp_path / "sample_contract.docx"
    doc = Document()
    doc.add_heading("Section 8: Service Levels", level=2)
    doc.add_paragraph("8.1 Uptime Commitment. Vendor guarantees 99.9% availability.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Availability Tier"
    table.rows[0].cells[1].text = "Service Credit"
    table.rows[1].cells[0].text = "< 99.9%"
    table.rows[1].cells[1].text = "10% monthly credit"
    doc.add_heading("Section 9: Confidentiality", level=2)
    doc.add_paragraph("9.1 Obligations. Both parties shall protect confidential data.")
    doc.save(str(docx_path))

    parser = DocumentParser()
    parsed = parser.parse(docx_path)

    clauses_by_id = {c.clause_id: c for c in parsed.clauses}
    assert "Section 8.1" in clauses_by_id
    assert "Section 9.1" in clauses_by_id
    # The SLA credit table must remain attached to Section 8.1 rather than pushed after Section 9.1
    assert "10% monthly credit" in clauses_by_id["Section 8.1"].text
    assert "10% monthly credit" not in clauses_by_id["Section 9.1"].text


def test_parser_unsupported_format(tmp_path):
    """Verifies defensive error handling on unsupported file extensions."""
    invalid_file = tmp_path / "contract.xyz"
    invalid_file.write_text("Unsupported dummy data", encoding="utf-8")

    parser = DocumentParser()
    with pytest.raises(ValueError, match="Unsupported file format"):
        parser.parse(invalid_file)


def test_parser_recursive_hierarchy_three_tiers(tmp_path):
    """Verifies recursive parsing of Addendum (L0) -> SOW (L1) -> MSA (L2) hierarchy."""
    msa_path = tmp_path / "base_msa.md"
    sow_path = tmp_path / "vendor_sow.md"
    addendum_path = tmp_path / "sow_addendum.md"

    msa_path.write_text(
        "## Section 14: General\n14.1 Governing Law. Governed by Delaware law.",
        encoding="utf-8",
    )
    sow_path.write_text(
        "## Section 3: Commercials\n3.2 Invoicing. Invoices payable Net 30 days.",
        encoding="utf-8",
    )
    addendum_path.write_text(
        "## Section 1: Amendments\n1.1 Payment Override. Notwithstanding Section 3.2, invoices are payable Net 60 days.",
        encoding="utf-8",
    )

    parser = DocumentParser()
    root = parser.parse_hierarchy(
        [addendum_path, sow_path, msa_path],
        labels=["Addendum", "SOW", "MSA"],
    )

    chain = root.iter_hierarchy()
    assert [node.document_label for node in chain] == ["Addendum", "SOW", "MSA"]

    clauses = root.resolve_hierarchical_clauses()
    qualified_ids = [c.qualified_id for c in clauses]
    assert "[Addendum (L0)] Section 1.1" in qualified_ids
    assert "[SOW (L1)] Section 3.2" in qualified_ids
    assert "[MSA (L2)] Section 14.1" in qualified_ids

    compiled_text = root.compile_hierarchical_text()
    assert "LEVEL 0 (PRIMARY / HIGHEST PRECEDENCE): Addendum" in compiled_text
    assert "LEVEL 1 (INTERMEDIATE GOVERNING AGREEMENT): SOW" in compiled_text
    assert "LEVEL 2 (BASE MASTER AGREEMENT): MSA" in compiled_text

    # Verify cycle guard raises ValueError if a parent pointer loops back
    assert root.parent_contract is not None
    assert root.parent_contract.parent_contract is not None
    root.parent_contract.parent_contract.parent_contract = root
    with pytest.raises(ValueError, match="Circular contract hierarchy"):
        root.iter_hierarchy()
