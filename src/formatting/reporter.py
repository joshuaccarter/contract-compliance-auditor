"""Reporting module to format and persist contract compliance audit results."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from src.audit.schemas import (
    ComplianceStatus,
    ConfidenceLevel,
    ContractAuditReport,
)


def _clean_md_cell(value: str) -> str:
    """Sanitizes pipe characters and newlines so Markdown table rows never break."""
    return " ".join(value.replace("|", "/").splitlines()).strip()


def _format_confidence_badge(confidence: ConfidenceLevel) -> str:
    """Returns a visual badge for the confidence level."""
    if confidence == ConfidenceLevel.HIGH:
        return "🟢 HIGH"
    if confidence == ConfidenceLevel.MEDIUM:
        return "🟡 MEDIUM"
    return "🔴 LOW"


def render_compliance_markdown(
    report: ContractAuditReport,
    output_path: str | Path,
) -> Path:
    """Renders a ContractAuditReport into a formatted Markdown intake sheet."""
    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)

    table_rows = []
    for r in report.results:
        badge = (
            "✅ MET"
            if r.status == ComplianceStatus.MET
            else (
                "⚠️ DEVIATION"
                if r.status == ComplianceStatus.DEVIATION
                else "❌ SILENT"
            )
        )
        conf_badge = _format_confidence_badge(r.confidence)
        cleaned_std = _clean_md_cell(r.policy_standard)
        cleaned_term = _clean_md_cell(r.identified_contract_term)
        cleaned_cite = _clean_md_cell(r.exact_citation)

        table_rows.append(
            f"| **{r.item_id}** | {r.item_name} | {cleaned_std} | "
            f"{cleaned_term} | {badge} (`{r.alignment_degree.value}`) | {conf_badge} | {cleaned_cite} |"
        )

    rows_text = "\n".join(table_rows)
    total = max(1, report.total_items_audited)

    content = f"""# 📋 Vendor Contract Intake & Compliance Assessment

**Contract Title:** {report.contract_title}  
**Vendor:** {report.vendor_name}  
**Overall Risk Assessment:** **{report.overall_risk_rating.value}**  

---

### Executive Compliance Scorecard

| Metric | Count | Percentage |
| :--- | :--- | :--- |
| **Total Checklist Items** | {report.total_items_audited} | 100% |
| **Fully Compliant (MET)** | {report.compliant_count} | {(report.compliant_count / total):.1%} |
| **Policy Deviations (DEVIATION)** | {report.deviation_count} | {(report.deviation_count / total):.1%} |
| **Missing / Unaddressed (SILENT)** | {report.silent_count} | {(report.silent_count / total):.1%} |
| **Confidence: HIGH / MEDIUM / LOW** | {report.high_confidence_count} / {report.medium_confidence_count} / {report.low_confidence_count} | {(report.high_confidence_count / total):.0%} High |

---

### Itemized Audit Review Matrix

| Item # | Focus Area | Policy Baseline Standard | Identified Vendor Term | Audit Verdict & Alignment | Confidence | Verbatim Citation / Section |
| :--- | :--- | :--- | :--- | :--- | :--- | :--- |
{rows_text}

---

### Actionable Findings & High-Exposure Items
"""

    deviations = [
        r
        for r in report.results
        if r.status in (ComplianceStatus.DEVIATION, ComplianceStatus.SILENT)
    ]
    if deviations:
        for d in deviations:
            status_text = (
                f"Policy Deviation — {d.alignment_degree.value}"
                if d.status == ComplianceStatus.DEVIATION
                else "Clause Missing (Silent)"
            )
            content += (
                f"\n* **[{d.item_id}] {d.item_name}** "
                f"({status_text} | Confidence: `{d.confidence.value}`):\n"
            )
            content += f"  * *Contract Term:* {d.identified_contract_term}\n"
            content += f"  * *Required Standard:* {d.policy_standard}\n"
            if d.deviation_analysis:
                content += f"  * *Exposure Risk:* {d.deviation_analysis}\n"
    else:
        content += "\n* No critical deviations or missing clauses identified. Contract aligns with enterprise standards.\n"

    content += "\n---\n*Report automatically compiled by Enterprise Contract Compliance Engine.*\n"

    dest.write_text(content, encoding="utf-8")
    return dest


def render_compliance_docx(
    report: ContractAuditReport,
    output_path: str | Path,
) -> Path:
    """Renders a ContractAuditReport into a formatted Microsoft Word document."""
    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Document Header
    title = doc.add_heading(
        "Vendor Contract Intake & Compliance Assessment", level=0
    )
    title.runs[0].font.size = Pt(18)

    meta_p = doc.add_paragraph()
    meta_p.add_run("Contract: ").bold = True
    meta_p.add_run(f"{report.contract_title}\n")
    meta_p.add_run("Vendor: ").bold = True
    meta_p.add_run(f"{report.vendor_name}\n")
    meta_p.add_run("Overall Risk Assessment: ").bold = True
    risk_run = meta_p.add_run(f"{report.overall_risk_rating.value}")
    risk_run.bold = True
    if report.overall_risk_rating.value in ("HIGH", "CRITICAL"):
        risk_run.font.color.rgb = RGBColor(180, 0, 0)
    else:
        risk_run.font.color.rgb = RGBColor(0, 120, 0)

    total = max(1, report.total_items_audited)

    # Executive Scorecard Table
    doc.add_heading("Executive Compliance Scorecard", level=2)
    summary_table = doc.add_table(rows=1, cols=3)
    summary_table.style = "Light Shading Accent 1"
    hdr = summary_table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Metric", "Count", "Percentage"

    stats = [
        ("Total Checklist Items", str(report.total_items_audited), "100%"),
        (
            "Fully Compliant (MET)",
            str(report.compliant_count),
            f"{(report.compliant_count / total):.1%}",
        ),
        (
            "Policy Deviations (DEVIATION)",
            str(report.deviation_count),
            f"{(report.deviation_count / total):.1%}",
        ),
        (
            "Missing / Unaddressed (SILENT)",
            str(report.silent_count),
            f"{(report.silent_count / total):.1%}",
        ),
        (
            "Confidence (HIGH / MEDIUM / LOW)",
            f"{report.high_confidence_count} / {report.medium_confidence_count} / {report.low_confidence_count}",
            f"{(report.high_confidence_count / total):.0%} High",
        ),
    ]
    for metric, count, pct in stats:
        row = summary_table.add_row().cells
        row[0].text, row[1].text, row[2].text = metric, count, pct

    # Detailed Audit Review Table
    doc.add_heading("Detailed Audit Review Matrix", level=2)
    detail_table = doc.add_table(rows=1, cols=6)
    detail_table.style = "Table Grid"
    hdr2 = detail_table.rows[0].cells
    hdr2[0].text = "ID"
    hdr2[1].text = "Focus Area"
    hdr2[2].text = "Identified Term"
    hdr2[3].text = "Verdict & Alignment"
    hdr2[4].text = "Confidence"
    hdr2[5].text = "Citation / Excerpt"

    for r in report.results:
        row = detail_table.add_row().cells
        row[0].text = r.item_id
        row[1].text = r.item_name
        row[2].text = r.identified_contract_term
        row[3].text = f"{r.status.value} ({r.alignment_degree.value})"
        row[4].text = r.confidence.value
        row[5].text = r.exact_citation

    # Exposure Items
    doc.add_heading("Actionable Findings & Deviations", level=2)
    deviations = [
        r
        for r in report.results
        if r.status in (ComplianceStatus.DEVIATION, ComplianceStatus.SILENT)
    ]
    if deviations:
        for d in deviations:
            tag = (
                f"Deviation: {d.alignment_degree.value}"
                if d.status == ComplianceStatus.DEVIATION
                else "Missing / Silent"
            )
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(
                f"[{d.item_id}] {d.item_name} ({tag} | Confidence: {d.confidence.value}): "
            ).bold = True
            p.add_run(
                f"Required: {d.policy_standard}. Found: {d.identified_contract_term}. "
            )
            if d.deviation_analysis:
                p.add_run(f"Exposure: {d.deviation_analysis}")
    else:
        doc.add_paragraph(
            "No critical deviations found. Terms align with standard policy."
        )

    doc.save(str(dest))
    return dest
